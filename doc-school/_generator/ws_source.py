# -*- coding: utf-8 -*-
"""
ws_source.py
============
Single source-of-truth loader for the workshop Excel file:

    doc-school/2607lts_nhóm 1_WebBshoes_Workshop1.xlsx

Parses the relevant sheets into clean, module-level Python data structures
that a docx generator can consume directly, without ever touching openpyxl
or the raw sheet layout again.

Run directly to print a validation report:

    PYTHONIOENCODING=utf-8 python ws_source.py
"""

import os
import unicodedata

import openpyxl

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------

_HERE = os.path.dirname(os.path.abspath(__file__))
XLSX_PATH = os.path.normpath(
    os.path.join(_HERE, "..", "2607lts_nhóm 1_WebBshoes_Workshop1.xlsx")
)

# ---------------------------------------------------------------------------
# Estimation constants (confirmed by user: "ED=18 meaning C=1")
# ---------------------------------------------------------------------------

ED_TOTAL = 18
ED_MAX = 36
C_DEFAULT = 1
ENV_FACTOR = ED_TOTAL / ED_MAX  # 0.5


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def _s(v):
    """Stringify + strip a cell value; None -> ''."""
    if v is None:
        return ""
    return str(v).strip()


def _is_x(v):
    """True if a cell marks an 'x' assignment (case-insensitive)."""
    return _s(v).lower() == "x"


def _to_int(v, default=0):
    if v is None or v == "":
        return default
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return default


def _to_float(v, default=0.0):
    if v is None or v == "":
        return default
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def _deaccent(s):
    """Strip Vietnamese diacritics -> plain ASCII letters."""
    s = s.replace("đ", "d").replace("Đ", "D")
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s


def _name_to_handle(name):
    """
    Derive the expected 'handle' for a Vietnamese full name, matching the
    convention observed in ws1.1 rows 14..19, e.g.:
        'Lưu Đình Bắc'        -> 'BacLD'
        'Hoàng Lê Bảo Ngọc'   -> 'NgocHLB'
        'Nguyễn Đình Dũng'    -> 'DungND'
        'Đỗ Anh Vũ'           -> 'VuDA'
    (given name last, capitalized + ascii; initials of the leading words,
    uppercase, appended)
    """
    parts = _deaccent(name).split()
    if not parts:
        return ""
    first = parts[-1]
    first = first[:1].upper() + first[1:].lower()
    initials = "".join(p[0].upper() for p in parts[:-1] if p)
    return first + initials


# ---------------------------------------------------------------------------
# Load workbook
# ---------------------------------------------------------------------------

_wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)


# ---------------------------------------------------------------------------
# Sheet ws1.1 -> TEAM + PROJECT
# ---------------------------------------------------------------------------

def _load_team_and_project():
    ws = _wb["ws1.1"]

    # --- raw team rows 5..10: B=role, C=name, D=phone ---
    raw_team = []
    for r in range(5, 11):
        role = _s(ws.cell(row=r, column=2).value)
        name = _s(ws.cell(row=r, column=3).value)
        phone = _s(ws.cell(row=r, column=4).value)
        if not name:
            continue  # skip blank rows (e.g. row 9)
        raw_team.append({"role": role, "name": name, "phone": phone})

    # --- display handle rows 14..19: B=handle, C=role, D=display ---
    handle_rows = []
    for r in range(14, 20):
        handle = _s(ws.cell(row=r, column=2).value)
        role = _s(ws.cell(row=r, column=3).value)
        display = _s(ws.cell(row=r, column=4).value)
        handle_rows.append({"handle": handle, "role": role, "display": display})

    # --- merge handle/display onto team members ---
    used_handle_rows = set()

    def _find_by_name_handle(member):
        expected = _name_to_handle(member["name"])
        for i, hr in enumerate(handle_rows):
            if i in used_handle_rows:
                continue
            if hr["handle"] and hr["handle"] == expected:
                return i
        return None

    def _find_by_unique_role(member):
        role_count_team = sum(1 for m in raw_team if m["role"] == member["role"])
        candidates = [
            i for i, hr in enumerate(handle_rows)
            if i not in used_handle_rows and hr["role"] == member["role"]
        ]
        if role_count_team == 1 and len(candidates) == 1:
            return candidates[0]
        return None

    team = []
    for member in raw_team:
        idx = _find_by_name_handle(member)
        if idx is None:
            idx = _find_by_unique_role(member)
        if idx is not None:
            used_handle_rows.add(idx)
            handle = handle_rows[idx]["handle"]
            display = handle_rows[idx]["display"]
        else:
            handle, display = "", ""
        team.append({
            "role": member["role"],
            "name": member["name"],
            "phone": member["phone"],
            "handle": handle,
            "display": display,
        })

    # --- project info: F/G rows 4..10 ---
    label_to_key = {
        "Tên dự án": "name",
        "Framework": "framework",
        "Sprint Duration": "sprint_duration",
        "Thời gian dự án": "time",
        "Cơ sở dữ liệu": "database",
        "Công nghệ": "tech",
        "Quy mô nhóm": "team_size",
    }
    # fallback: hardcoded row order if a label doesn't match exactly
    row_to_key_fallback = {
        4: "name", 5: "framework", 6: "sprint_duration", 7: "time",
        8: "database", 9: "tech", 10: "team_size",
    }
    project = {}
    for r in range(4, 11):
        label = _s(ws.cell(row=r, column=6).value)
        value = _s(ws.cell(row=r, column=7).value)
        key = label_to_key.get(label, row_to_key_fallback.get(r))
        if key:
            project[key] = value

    return team, project


TEAM, PROJECT = _load_team_and_project()


# ---------------------------------------------------------------------------
# Sheet 'ws 1.2' -> WORK_ASSIGN
# ---------------------------------------------------------------------------

def _load_work_assign():
    ws = _wb["ws 1.2"]
    out = []
    for r in range(3, 14):
        task = _s(ws.cell(row=r, column=2).value)
        if not task:
            continue
        out.append({
            "stt": _to_int(ws.cell(row=r, column=1).value, default=len(out) + 1),
            "task": task,
            "po": _is_x(ws.cell(row=r, column=3).value),
            "sm": _is_x(ws.cell(row=r, column=4).value),
            "team": _is_x(ws.cell(row=r, column=5).value),
            "pm": _is_x(ws.cell(row=r, column=6).value),
        })
    return out


WORK_ASSIGN = _load_work_assign()


# ---------------------------------------------------------------------------
# Sheet '1.0 Use case' -> USECASES
# ---------------------------------------------------------------------------

def _load_usecases():
    ws = _wb["1.0 Use case"]
    out = []
    for r in range(4, 11):
        group = _s(ws.cell(row=r, column=6).value)
        if not group:
            continue
        uses_raw = _s(ws.cell(row=r, column=7).value)
        tables_raw = _s(ws.cell(row=r, column=8).value)
        pic = _s(ws.cell(row=r, column=4).value)
        uses = [ln.strip() for ln in uses_raw.split("\n") if ln.strip()]
        tables = [ln.strip() for ln in tables_raw.split("\n") if ln.strip()]
        out.append({"group": group, "uses": uses, "tables": tables, "pic": pic})
    return out


USECASES = _load_usecases()


# ---------------------------------------------------------------------------
# Sheet 'ws1.3.1. Product Backlog' -> RQS (40 requirements)
# ---------------------------------------------------------------------------

def _load_rqs():
    ws = _wb["ws1.3.1. Product Backlog"]
    out = []
    for r in range(5, 45):  # rows 5..44 inclusive (RQ-1..RQ-40)
        rq_id = _s(ws.cell(row=r, column=1).value)
        if not rq_id or rq_id == "RQ-0":
            continue
        role = _s(ws.cell(row=r, column=2).value)
        goal = _s(ws.cell(row=r, column=3).value)
        so_that = _s(ws.cell(row=r, column=4).value)
        priority = _to_int(ws.cell(row=r, column=5).value, default=0)
        bv = _s(ws.cell(row=r, column=6).value)
        sp = _to_float(ws.cell(row=r, column=7).value, default=0.0)
        state = _s(ws.cell(row=r, column=8).value)
        out.append({
            "id": rq_id, "role": role, "goal": goal, "so_that": so_that,
            "priority": priority, "bv": bv, "sp": sp, "state": state,
        })
    return out


RQS = _load_rqs()
RQ_BY_ID = {r["id"]: r for r in RQS}


# ---------------------------------------------------------------------------
# Sheet 'ws1.3.2. Release Backlog' -> PBS (13 groups) + STORY_SPRINT
# ---------------------------------------------------------------------------

def _load_release_backlog():
    ws = _wb["ws1.3.2. Release Backlog"]

    pb_order = []          # preserves first-seen order of PB ids
    pb_name = {}            # pb_id -> group name
    pb_stories = {}         # pb_id -> [rq ids]
    pb_sprints = {}         # pb_id -> set of sprint ints
    story_sprint = {}       # rq_id -> sprint int

    for r in range(5, 45):  # rows 5..44
        pb_id = _s(ws.cell(row=r, column=1).value)
        if not pb_id:
            continue
        name = _s(ws.cell(row=r, column=3).value)
        story_id = _s(ws.cell(row=r, column=7).value)
        sprint = _to_int(ws.cell(row=r, column=10).value, default=None)

        if pb_id not in pb_stories:
            pb_order.append(pb_id)
            pb_stories[pb_id] = []
            pb_sprints[pb_id] = set()

        if name:
            pb_name[pb_id] = name

        if story_id:
            pb_stories[pb_id].append(story_id)
            if sprint is not None:
                pb_sprints[pb_id].add(sprint)
                story_sprint[story_id] = sprint

    pbs = []
    for pb_id in pb_order:
        sprints_sorted = sorted(pb_sprints[pb_id])
        pbs.append({
            "id": pb_id,
            "name": pb_name.get(pb_id, ""),
            "stories": pb_stories[pb_id],
            "sprint": sprints_sorted[0] if sprints_sorted else None,
            "sprints": sprints_sorted,
        })

    return pbs, story_sprint


PBS, STORY_SPRINT = _load_release_backlog()


# ---------------------------------------------------------------------------
# Sheet 'ws1.3.4 Sprint Planning' -> SPRINTS, DEFENSE, TOTALS
# ---------------------------------------------------------------------------

def _fmt_date(v):
    if v is None:
        return ""
    try:
        return v.strftime("%d/%m/%Y")
    except AttributeError:
        return _s(v)


def _load_sprint_planning():
    ws = _wb["ws1.3.4 Sprint Planning"]

    sprints = []
    for r in range(3, 9):  # rows 3..8 -> sprint 1..6
        num = _to_int(ws.cell(row=r, column=1).value, default=r - 2)
        raw_b = _s(ws.cell(row=r, column=2).value)
        lines = raw_b.split("\n")
        first_line = lines[0] if lines else ""
        extra_lines = [ln.strip() for ln in lines[1:] if ln.strip()]

        if " - " in first_line:
            name_part, theme_part = first_line.split(" - ", 1)
        else:
            name_part, theme_part = first_line, ""
        name = name_part.strip()
        theme_bits = [b.strip() for b in [theme_part] + extra_lines if b.strip()]
        theme = "\n".join(theme_bits)

        sprints.append({
            "num": num,
            "name": name,
            "theme": theme,
            "pic": _s(ws.cell(row=r, column=3).value),
            "goal": _s(ws.cell(row=r, column=4).value),
            "start": _fmt_date(ws.cell(row=r, column=5).value),
            "end": _fmt_date(ws.cell(row=r, column=6).value),
            "us": _to_int(ws.cell(row=r, column=7).value, default=0),
            "pps": _to_float(ws.cell(row=r, column=8).value, default=0.0),
        })

    # row 9 = 'Bảo vệ' + TỔNG
    defense = {
        "label": _s(ws.cell(row=9, column=2).value),
        "pic": _s(ws.cell(row=9, column=3).value),
        "goal": _s(ws.cell(row=9, column=4).value),
    }
    totals = {
        "us": _to_int(ws.cell(row=9, column=7).value, default=0),
        "pps": _to_float(ws.cell(row=9, column=8).value, default=0.0),
    }

    return sprints, defense, totals


SPRINTS, DEFENSE, TOTALS = _load_sprint_planning()


# ---------------------------------------------------------------------------
# Derived Sprint Backlog tasks (NOT parsed from the messy ws1.3.3 sheet;
# reproduced deterministically from PBS + RQS + STORY_SPRINT)
# ---------------------------------------------------------------------------

def _load_tasks():
    tasks = []
    n = 0

    def _next_id():
        nonlocal n
        n += 1
        return "T-%d" % n

    for pb in PBS:
        group_name = pb["name"]
        group_sprint = pb["sprint"]

        tasks.append({
            "id": _next_id(),
            "task": "Vẽ use case %s" % group_name,
            "desc": "",
            "story": "",
            "backlog": pb["id"],
            "sprint": group_sprint,
            "est": 1,
            "who": "DungND",
        })
        tasks.append({
            "id": _next_id(),
            "task": "Phân tích và thiết kế sơ đồ khối chức năng %s" % group_name,
            "desc": "",
            "story": "",
            "backlog": pb["id"],
            "sprint": group_sprint,
            "est": 1,
            "who": "DungND",
        })

        for rq_id in pb["stories"]:
            rq = RQ_BY_ID.get(rq_id, {})
            goal_text = rq.get("goal", "")
            tasks.append({
                "id": _next_id(),
                "task": "Code: %s" % goal_text,
                "desc": goal_text,
                "story": rq_id,
                "backlog": pb["id"],
                "sprint": STORY_SPRINT.get(rq_id, group_sprint),
                "est": 4,
                "who": "BacLD",
            })

        tasks.append({
            "id": _next_id(),
            "task": "Tích hợp tính năng",
            "desc": "",
            "story": "",
            "backlog": pb["id"],
            "sprint": group_sprint,
            "est": 2,
            "who": "VuDA",
        })
        tasks.append({
            "id": _next_id(),
            "task": "Triển khai kiểm thử đơn vị, tích hợp và hệ thống cho %s" % group_name,
            "desc": "",
            "story": "",
            "backlog": pb["id"],
            "sprint": group_sprint,
            "est": 8,
            "who": "NgocHLB",
        })

    return tasks


TASKS = _load_tasks()


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    print("XLSX_PATH:", XLSX_PATH)
    print()

    print("len(TEAM) =", len(TEAM))
    for m in TEAM:
        print("   ", m)
    print()

    print("PROJECT =", PROJECT)
    print()

    print("len(WORK_ASSIGN) =", len(WORK_ASSIGN))
    print()

    print("len(USECASES) =", len(USECASES))
    for u in USECASES:
        print("   ", u["group"], "-> pic:", u["pic"], "| uses:", len(u["uses"]), "| tables:", len(u["tables"]))
    print()

    n_new = sum(1 for r in RQS if r["state"] == "New")
    n_removed = sum(1 for r in RQS if r["state"] == "Removed")
    print("len(RQS) =", len(RQS), "(New=%d, Removed=%d)" % (n_new, n_removed))
    print()

    print("len(PBS) =", len(PBS))
    for pb in PBS:
        print("   %-6s %-40s stories=%-3d sprint=%s sprints=%s" % (
            pb["id"], pb["name"], len(pb["stories"]), pb["sprint"], pb["sprints"]
        ))
    print()

    print("len(SPRINTS) =", len(SPRINTS))
    for s in SPRINTS:
        print("    Sprint %d | %s -> %s | US=%s | PPS=%s | %s | theme=%r" % (
            s["num"], s["start"], s["end"], s["us"], s["pps"], s["pic"], s["theme"]
        ))
    print("DEFENSE =", DEFENSE)
    print("TOTALS =", TOTALS)
    print()

    print("len(TASKS) =", len(TASKS))
    print("First 6 tasks:")
    for t in TASKS[:6]:
        print("   ", t)
    print("Last 3 tasks:")
    for t in TASKS[-3:]:
        print("   ", t)
    print()

    assert len(RQS) == 40, "expected 40 RQS, got %d" % len(RQS)
    assert n_new == 24, "expected 24 New RQS, got %d" % n_new
    assert len(PBS) == 13, "expected 13 PBS, got %d" % len(PBS)
    assert len(SPRINTS) == 6, "expected 6 SPRINTS, got %d" % len(SPRINTS)

    print("ALL ASSERTS PASSED")
