# -*- coding: utf-8 -*-
"""Đọc ngược file docx đã sinh để kiểm chứng định dạng, thay vì tin generator."""
import sys, os, re
from docx import Document

LA_MA = re.compile(r'^\s*(I|II|III|IV|V|VI|VII|VIII|IX|X)[.\s]')
bad = []

for path in sys.argv[1:]:
    d = Document(path)
    name = os.path.basename(path.replace('\\', '/'))
    print("=" * 70)
    print(name)

    # 1. gạch ngang dài
    dashes = 0
    for p in d.paragraphs:
        for ch in ('—', '–'):
            dashes += p.text.count(ch)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for ch in ('—', '–'):
                    dashes += c.text.count(ch)
    print("  gach ngang dai (em/en):", dashes)
    if dashes:
        bad.append("%s: con %d gach ngang dai" % (name, dashes))

    # 2. font Normal
    st = d.styles['Normal']
    size = st.font.size.pt if st.font.size else None
    ls = st.paragraph_format.line_spacing
    print("  Normal:", st.font.name, "| size:", size, "| line spacing:", ls)
    if st.font.name != 'Times New Roman':
        bad.append("%s: Normal khong phai Times New Roman" % name)
    if size != 13:
        bad.append("%s: Normal khong phai 13pt" % name)
    if not ls or not (1.15 <= float(ls) <= 1.5):
        bad.append("%s: line spacing %s ngoai khoang 1.15-1.5" % (name, ls))

    # 3. chương: số Ả Rập, không La Mã
    chuongs = [p.text for p in d.paragraphs if p.text.strip().startswith('CHƯƠNG')]
    print("  so chuong:", len(chuongs))
    for c in chuongs:
        print("     ", c)
    for p in d.paragraphs:
        if p.style.name == 'Heading 1' and LA_MA.match(p.text):
            bad.append("%s: con so La Ma: %s" % (name, p.text))

    # 4. caption: Heading 5 = tên hình, Heading 6 = tên bảng
    h5 = [p.text for p in d.paragraphs if p.style.name == 'Heading 5']
    h6 = [p.text for p in d.paragraphs if p.style.name == 'Heading 6']
    print("  Heading 5 (ten hinh):", len(h5))
    print("  Heading 6 (ten bang):", len(h6))
    for x in h6[:3]:
        print("      ", x)
    for x in h6:
        if not re.match(r'^Bảng \d+\.\d+: ', x):
            bad.append("%s: caption bang sai dinh dang: %s" % (name, x))

    # 5. TITLE: đậm, canh giữa, màu thương hiệu
    p0 = d.paragraphs[0]
    for r in p0.runs:
        print("  TITLE: bold=%s color=%s size=%s align=%s"
              % (r.font.bold, r.font.color.rgb, r.font.size.pt if r.font.size else None, p0.alignment))
        if not r.font.bold:
            bad.append("%s: title khong dam" % name)
        if str(r.font.color.rgb) != '0B895A':
            bad.append("%s: title sai mau thuong hieu" % name)
    if p0.alignment != 1:  # CENTER
        bad.append("%s: title khong canh giua" % name)

    # 6. H2/H3 lùi đầu dòng
    h2s = [p for p in d.paragraphs if p.style.name == 'Heading 2']
    h3s = [p for p in d.paragraphs if p.style.name == 'Heading 3']
    if h2s:
        ind = h2s[0].paragraph_format.left_indent
        print("  H2 indent:", ind.cm if ind else None, "| vd:", h2s[0].text[:40])
        if not ind or ind.cm <= 0:
            bad.append("%s: H2 khong lui dau dong" % name)
    if h3s:
        ind3 = h3s[0].paragraph_format.left_indent
        it = all(r.font.italic for r in h3s[0].runs)
        print("  H3 indent:", ind3.cm if ind3 else None, "| italic:", it, "| vd:", h3s[0].text[:40])
        if not ind3 or ind3.cm <= 0:
            bad.append("%s: H3 khong lui dau dong" % name)
        if not it:
            bad.append("%s: H3 khong in nghieng" % name)

    print("  KET LUAN:", any(p.text.strip() == 'KẾT LUẬN' for p in d.paragraphs))
    print("  so bang:", len(d.tables))

print("=" * 70)
if bad:
    print("LOI:")
    for b in bad:
        print("  -", b)
    sys.exit(1)
print("Tat ca kiem tra dinh dang: PASS")
