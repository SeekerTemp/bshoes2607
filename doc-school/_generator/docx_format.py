# -*- coding: utf-8 -*-
"""
Bộ khuôn định dạng Word theo yêu cầu người dùng (2026-07-17).

  Font       : Times New Roman, cỡ 13, giãn dòng 1.15 đến 1.5
  TITLE      : đậm, canh giữa, màu thương hiệu
  CHƯƠNG n   : H1, in đậm, màu logo, KHÔNG dùng số La Mã
  n.n        : H2, đậm, lùi 2 ký tự
  n.n.n      : H3, đậm, nghiêng, có màu, lùi 4 đến 5 ký tự
  Tên hình   : Heading 5   |   Tên bảng : Heading 6
  Đánh số    : theo chương, ví dụ 3.7 là hình thứ 7 của chương 3
  Chú thích bảng đặt TRÊN bảng, chú thích hình đặt DƯỚI hình

Tuyệt đối không dùng dấu gạch ngang dài (em dash và en dash) trong nội dung.
Hàm kiem_tra_gach_ngang() chặn ngay lúc sinh file nếu lỡ lọt.
"""
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

FONT = "Times New Roman"
SIZE = Pt(13)
GIAN_DONG = 1.3                       # nằm giữa khoảng 1.15 đến 1.5 được yêu cầu
GREEN = RGBColor(0x0B, 0x89, 0x5A)    # màu logo BShoes
GACH_NGANG_DAI = ("—", "–")  # em dash, en dash


def kiem_tra_gach_ngang(*texts):
    """Chặn dấu gạch ngang dài ngay tại nguồn, thay vì phát hiện sau khi mở Word."""
    for t in texts:
        if t is None:
            continue
        for d in GACH_NGANG_DAI:
            if d in str(t):
                raise ValueError("Còn dấu gạch ngang dài %r trong: %r" % (d, t))


def _font_run(run, size=SIZE, bold=None, italic=None, color=None):
    run.font.name = FONT
    run.font.size = size
    # eastAsia phải set riêng, nếu không Word vẫn nhảy về font mặc định với chữ có dấu
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def _style_para(p, indent_cm=0.0, align=None, spacing=GIAN_DONG):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = spacing
    pf.space_after = Pt(6)
    pf.left_indent = Cm(indent_cm)
    if align is not None:
        p.alignment = align
    return p


def setup(doc):
    """Đặt font/giãn dòng cho toàn tài liệu qua style Normal và các style heading."""
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = SIZE
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = GIAN_DONG
    pf.space_after = Pt(6)
    # Heading mặc định của Word là Calibri Light màu xanh dương, phải ghi đè hết.
    for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Heading 6", "Title"):
        try:
            s = doc.styles[name]
        except KeyError:
            continue
        s.font.name = FONT
        s.font.color.rgb = GREEN
        s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
    return doc


# ---------------------------------------------------------------- tiêu đề

def title(doc, text, phu=None):
    """TITLE: đậm, canh giữa, màu thương hiệu."""
    kiem_tra_gach_ngang(text, phu)
    p = doc.add_paragraph()
    _style_para(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    _font_run(p.add_run(text.upper()), size=Pt(18), bold=True, color=GREEN)
    if phu:
        p2 = doc.add_paragraph()
        _style_para(p2, align=WD_ALIGN_PARAGRAPH.CENTER)
        _font_run(p2.add_run(phu), italic=True)
    return p


class Chuong:
    """
    Một chương. Giữ bộ đếm hình và bảng để đánh số theo dạng <chương>.<thứ tự>.
    """

    def __init__(self, doc, so, ten):
        self.doc = doc
        self.so = so
        self._hinh = 0
        self._bang = 0
        kiem_tra_gach_ngang(ten)
        p = doc.add_heading(level=1)
        _style_para(p)
        # "CHƯƠNG 1: TÊN CHƯƠNG", số Ả Rập, không La Mã
        _font_run(p.add_run("CHƯƠNG %d: %s" % (so, ten.upper())), size=Pt(15), bold=True, color=GREEN)

    def h2(self, muc, text):
        """1.1. Tiêu đề 2: đậm, lùi 2 ký tự."""
        kiem_tra_gach_ngang(text)
        p = self.doc.add_heading(level=2)
        _style_para(p, indent_cm=0.5)          # 2 ký tự Times New Roman 13 xấp xỉ 0.5cm
        _font_run(p.add_run("%d.%d. %s" % (self.so, muc, text)), size=Pt(14), bold=True, color=GREEN)
        return p

    def h3(self, muc, submuc, text):
        """1.1.1. Tiêu đề 3: đậm, nghiêng, có màu, lùi 4 đến 5 ký tự."""
        kiem_tra_gach_ngang(text)
        p = self.doc.add_heading(level=3)
        _style_para(p, indent_cm=1.1)          # xấp xỉ 4 đến 5 ký tự
        _font_run(p.add_run("%d.%d.%d. %s" % (self.so, muc, submuc, text)),
                  size=Pt(13), bold=True, italic=True, color=GREEN)
        return p

    # ------------------------------------------------------------ nội dung

    def para(self, text):
        kiem_tra_gach_ngang(text)
        p = self.doc.add_paragraph()
        _style_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        _font_run(p.add_run(text))
        return p

    def bullet_head(self, text):
        """Ý lớn."""
        kiem_tra_gach_ngang(text)
        p = self.doc.add_paragraph(style="List Bullet")
        _style_para(p, indent_cm=0.8)
        for r in p.runs:
            _font_run(r)
        _font_run(p.add_run(text), bold=True)
        return p

    def bullet_item(self, text):
        """Ý nhỏ, gạch đầu dòng (dùng gạch nối ngắn, không phải gạch ngang dài)."""
        kiem_tra_gach_ngang(text)
        p = self.doc.add_paragraph()
        _style_para(p, indent_cm=1.6)
        _font_run(p.add_run("- " + text))
        return p

    def so_thu_tu(self, text, i):
        kiem_tra_gach_ngang(text)
        p = self.doc.add_paragraph()
        _style_para(p, indent_cm=0.8)
        _font_run(p.add_run("%d. " % i), bold=True)
        _font_run(p.add_run(text))
        return p

    # ------------------------------------------------------------ bảng, hình

    def bang(self, ten, headers, rows, widths=None):
        """
        Chú thích bảng nằm TRÊN bảng, đánh số <chương>.<thứ tự>, style Heading 6.
        """
        self._bang += 1
        kiem_tra_gach_ngang(ten, *headers)
        cap = self.doc.add_heading(level=6)
        _style_para(cap, align=WD_ALIGN_PARAGRAPH.CENTER)
        _font_run(cap.add_run("Bảng %d.%d: %s" % (self.so, self._bang, ten)),
                  size=Pt(12), bold=True, italic=True, color=GREEN)

        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ""
            _style_para(c.paragraphs[0], align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.15)
            _font_run(c.paragraphs[0].add_run(str(h)), size=Pt(11), bold=True, color=GREEN)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                kiem_tra_gach_ngang(v)
                cells[i].text = ""
                _style_para(cells[i].paragraphs[0], spacing=1.15)
                _font_run(cells[i].paragraphs[0].add_run("" if v is None else str(v)), size=Pt(11))
        if widths:
            for r in t.rows:
                for i, w in enumerate(widths):
                    r.cells[i].width = Cm(w)
        self.doc.add_paragraph()
        return t

    def hinh(self, ten, duong_dan=None, rong_cm=14):
        """
        Chú thích hình nằm DƯỚI hình, đánh số <chương>.<thứ tự>, style Heading 5.
        """
        self._hinh += 1
        kiem_tra_gach_ngang(ten)
        if duong_dan:
            p = self.doc.add_paragraph()
            _style_para(p, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.add_run().add_picture(duong_dan, width=Cm(rong_cm))
        cap = self.doc.add_heading(level=5)
        _style_para(cap, align=WD_ALIGN_PARAGRAPH.CENTER)
        _font_run(cap.add_run("Hình %d.%d: %s" % (self.so, self._hinh, ten)),
                  size=Pt(12), bold=True, italic=True, color=GREEN)
        return cap


def ket_luan(doc, cac_doan):
    """Phần Kết luận cuối tài liệu, đánh như một chương nhưng không có số."""
    p = doc.add_heading(level=1)
    _style_para(p)
    _font_run(p.add_run("KẾT LUẬN"), size=Pt(15), bold=True, color=GREEN)
    for d in cac_doan:
        kiem_tra_gach_ngang(d)
        q = doc.add_paragraph()
        _style_para(q, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        _font_run(q.add_run(d))
