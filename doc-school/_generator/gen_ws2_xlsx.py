# -*- coding: utf-8 -*-
"""Sinh WORKSHOP_2_BShoes.xlsx: Product Backlog (RQ), Release Backlog (PB),
Sprint Planning, Sprint Backlog (tách từng sprint), ED, Ước lượng Story.

Bám cột theo file mẫu doc-school-request (3 sheet Product/Release/Sprint Backlog).
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from bshoes_data import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
GREEN = "0B895A"; LIGHT = "E7F4EF"; GREY = "EEF1F4"; BAND = "D9EAE2"
hdr_font = Font(bold=True, color="FFFFFF", size=11)
hdr_fill = PatternFill("solid", fgColor=GREEN)
title_font = Font(bold=True, size=14, color=GREEN)
band_font = Font(bold=True, color=GREEN, size=11)
band_fill = PatternFill("solid", fgColor=BAND)
bold = Font(bold=True)
thin = Side(style="thin", color="D5DBE2")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
ctr = Alignment(horizontal="center", vertical="center")

wb = Workbook()

NPB = len(PBS)
NRQ = len(RQS)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def goal_of(user_story):
    """Tach phan 'I want to' khoi cau 'La <role>, toi muon <goal>'."""
    return user_story.split("tôi muốn ", 1)[1] if "tôi muốn " in user_story else user_story


# ============ 0. Tổng quan ============
ws = wb.active; ws.title = "0. Tổng quan"
ws["A1"] = PROJECT; ws["A1"].font = title_font
ws["A2"] = "WORKSHOP 2: Product Backlog, Release Backlog, Sprint Backlog & Ước lượng Story"; ws["A2"].font = bold
ws["A3"] = "Phiên bản: " + VERSION
ws["A4"] = "Hạn hoàn thành: " + DEADLINE
r = 6
ws.cell(row=r, column=1, value="Cấu trúc tài liệu (3 cấp, đánh số theo Sprint)").font = bold; r += 1
for line in ["Cấp 1, RQ (Request): yêu cầu từ actor, dạng As a / I want / So that -> Product Backlog",
             "Cấp 2, PB (Product Backlog Item): user story bóc từ RQ, có Priority/Sprint/Story Point -> Release Backlog",
             "Cấp 3, Task: mỗi PB chia thành nhiều Task -> Sprint Backlog (tách từng Sprint)"]:
    ws.cell(row=r, column=1, value=line); r += 1
r += 1
ws.cell(row=r, column=1, value="Công thức ước lượng (theo tài liệu 'Ước lượng số lượng Story')").font = bold; r += 1
for line in ["UP  (Điểm chưa hiệu chỉnh) = Loại tương tác + Quy tắc nghiệp vụ + Số thực thể + Thao tác dữ liệu   (mỗi mục 1..3)",
             "AP  (Điểm đã hiệu chỉnh)   = UP × C   (C = hệ số nhân)",
             "PPS (Điểm cho mỗi Story)   = (AP × ED) / 36",
             "ED  = tổng 18 yếu tố môi trường, mỗi yếu tố 0 hoặc 2  ->  tối đa 36"]:
    ws.cell(row=r, column=1, value=line); r += 1
r += 1
ws.cell(row=r, column=1, value="Kết quả tính cho nhóm").font = bold; r += 1
ws.cell(row=r, column=1, value="ED của nhóm"); ws.cell(row=r, column=2, value="='5. ED (Môi trường)'!D23"); r += 1
ws.cell(row=r, column=1, value="Hệ số môi trường ED/36"); ws.cell(row=r, column=2, value="='5. ED (Môi trường)'!D24"); r += 1
ws.cell(row=r, column=1, value="C (hệ số nhân)"); ws.cell(row=r, column=2, value="='6. Ước lượng Story'!C2"); r += 2
ws.cell(row=r, column=1, value="Nhóm thực hiện").font = bold; r += 1
ws.cell(row=r, column=1, value="Vai trò").font = bold
ws.cell(row=r, column=2, value="Thành viên").font = bold
ws.cell(row=r, column=3, value="Trách nhiệm").font = bold; r += 1
for code, name, resp in TEAM:
    ws.cell(row=r, column=1, value=code); ws.cell(row=r, column=2, value=name)
    ws.cell(row=r, column=3, value=resp); r += 1
widths(ws, [22, 26, 96])

# ============ 1. Product Backlog (cấp Request) ============
ws = wb.create_sheet("1. Product Backlog")
ws["A1"] = "Product Backlog cấp Request (RQ): As a / I want / So that"; ws["A1"].font = title_font
cols = ["ID", "As a/an [role]", "I want to [goal]", "So that [reason]", "Priority",
        "Business Value", "Acceptance Criteria", "State"]
ws.append([]); ws.append(cols); style_header(ws, 3, len(cols))
r = 4
for rq in RQS:
    ws.cell(row=r, column=1, value=rq["id"])
    ws.cell(row=r, column=2, value="Là " + rq["role"])
    ws.cell(row=r, column=3, value=rq["goal"])
    ws.cell(row=r, column=4, value=rq["so_that"])
    ws.cell(row=r, column=5, value=rq["priority"])
    ws.cell(row=r, column=6, value=rq["bv"])
    ws.cell(row=r, column=7, value="Chức năng hoạt động đúng và được PO nghiệm thu")
    ws.cell(row=r, column=8, value="New")
    for c in range(1, len(cols) + 1):
        ws.cell(row=r, column=c).border = border; ws.cell(row=r, column=c).alignment = wrap
    ws.cell(row=r, column=5).alignment = ctr
    r += 1
widths(ws, [9, 20, 46, 44, 9, 14, 40, 9])

# ============ 2. Release Backlog (cấp Product Backlog Item) ============
ws = wb.create_sheet("2. Release Backlog")
ws["A1"] = "Release Backlog: user story theo Backlog (RQ) và Sprint"; ws["A1"].font = title_font
cols = ["Backlog ID", "Backlog", "As a/an [role]", "I want to [goal]", "So that [reason]",
        "Story ID", "Priority", "Business Value", "Sprint#", "State", "Story Points"]
ws.append([]); ws.append(cols); style_header(ws, 3, len(cols))
r = 4
rq_ten = {x["id"]: x["ten"] for x in RQS}
for rq in RQS:
    for pb in [p for p in PBS if p["rq"] == rq["id"]]:
        ws.cell(row=r, column=1, value=rq["id"])
        ws.cell(row=r, column=2, value=rq["ten"])
        ws.cell(row=r, column=3, value="Là " + pb["role_name"])
        ws.cell(row=r, column=4, value=goal_of(pb["user_story"]))
        ws.cell(row=r, column=5, value=pb["mo_ta"])
        ws.cell(row=r, column=6, value=pb["id"])
        ws.cell(row=r, column=7, value=pb["priority"])
        ws.cell(row=r, column=8, value=pb["bv"])
        ws.cell(row=r, column=9, value=pb["sprint"])
        ws.cell(row=r, column=10, value="New")
        ws.cell(row=r, column=11, value="=VLOOKUP(F%d,'6. Ước lượng Story'!$A$5:$L$%d,12,FALSE)" % (r, 4 + NPB))
        for c in range(1, len(cols) + 1):
            ws.cell(row=r, column=c).border = border; ws.cell(row=r, column=c).alignment = wrap
        for c in (7, 9, 11):
            ws.cell(row=r, column=c).alignment = ctr
        r += 1
rb_first, rb_last = 4, r - 1
ws.cell(row=r, column=6, value="TỔNG").font = bold
ws.cell(row=r, column=11, value="=SUM(K%d:K%d)" % (rb_first, rb_last)).font = bold
widths(ws, [11, 22, 16, 40, 40, 9, 9, 14, 9, 9, 12])

# ============ 3. Sprint Planning ============
ws = wb.create_sheet("3. Sprint Planning")
ws["A1"] = "Sprint Duration: 2 tuần"; ws["A1"].font = bold
cols = ["Sprint", "Mục tiêu Sprint (Sprint Goal)", "Số User Story", "Tổng Story Points", "Tổng giờ task"]
ws.append([]); ws.append(cols); style_header(ws, 3, len(cols))
r = 4
for sp, goal in SPRINTS:
    ws.cell(row=r, column=1, value="Sprint %d" % sp)
    ws.cell(row=r, column=2, value=goal)
    ws.cell(row=r, column=3, value="=COUNTIF('2. Release Backlog'!$I$%d:$I$%d,%d)" % (rb_first, rb_last, sp))
    ws.cell(row=r, column=4, value="=SUMIF('2. Release Backlog'!$I$%d:$I$%d,%d,'2. Release Backlog'!$K$%d:$K$%d)" % (rb_first, rb_last, sp, rb_first, rb_last))
    ws.cell(row=r, column=5, value="=SUMIF('4. Sprint Backlog'!$F$2:$F$600,%d,'4. Sprint Backlog'!$G$2:$G$600)" % sp)
    for c in range(1, len(cols) + 1):
        ws.cell(row=r, column=c).border = border; ws.cell(row=r, column=c).alignment = wrap
    r += 1
ws.cell(row=r, column=2, value="TỔNG").font = bold
ws.cell(row=r, column=3, value="=SUM(C4:C%d)" % (r - 1)).font = bold
ws.cell(row=r, column=4, value="=SUM(D4:D%d)" % (r - 1)).font = bold
ws.cell(row=r, column=5, value="=SUM(E4:E%d)" % (r - 1)).font = bold
widths(ws, [11, 70, 14, 17, 14])

# ============ 4. Sprint Backlog (tách từng Sprint) ============
ws = wb.create_sheet("4. Sprint Backlog")
cols = ["Task ID", "Task", "Mô tả công việc", "Story ID", "Backlog ID", "Sprint#",
        "Estimate (giờ)", "Phụ trách", "Trạng thái"]
ws.append(cols); style_header(ws, 1, len(cols))
r = 2
for sp, goal in SPRINTS:
    ws.cell(row=r, column=1, value="Sprint %d: %s" % (sp, goal))
    for c in range(1, len(cols) + 1):
        ws.cell(row=r, column=c).fill = band_fill; ws.cell(row=r, column=c).border = border
    ws.cell(row=r, column=1).font = band_font
    try:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
    except Exception:
        pass
    r += 1
    for t in [t for t in TASKS if t["sprint"] == sp]:
        ws.cell(row=r, column=1, value=t["id"])
        ws.cell(row=r, column=2, value=t["task"])
        ws.cell(row=r, column=3, value=t["mo_ta"])
        ws.cell(row=r, column=4, value=t["story"])
        ws.cell(row=r, column=5, value=t["backlog"])
        ws.cell(row=r, column=6, value=t["sprint"])
        ws.cell(row=r, column=7, value=t["est"])
        ws.cell(row=r, column=8, value=t["who"])
        ws.cell(row=r, column=9, value="To Do")
        for c in range(1, len(cols) + 1):
            ws.cell(row=r, column=c).border = border; ws.cell(row=r, column=c).alignment = wrap
        for c in (4, 5, 6, 7):
            ws.cell(row=r, column=c).alignment = ctr
        r += 1
ws.cell(row=r, column=3, value="TỔNG GIỜ").font = bold
ws.cell(row=r, column=7, value="=SUM(G2:G%d)" % (r - 1)).font = bold
widths(ws, [9, 38, 44, 10, 11, 8, 13, 11, 11])

# ============ 5. ED (Môi trường) ============
ws = wb.create_sheet("5. ED (Môi trường)")
ws["A1"] = "ED: Đánh giá môi trường (Environment Degree)"; ws["A1"].font = title_font
ws["A2"] = "Mỗi yếu tố chấm 0 hoặc 2 (2 = thuận lợi cho nhóm). Tổng tối đa = 18 yếu tố × 2 = 36."
cols = ["Khía cạnh", "Yếu tố", "Khoảng giá trị", "Giá trị (0/2)", "Lý do chấm điểm của nhóm"]
ws.append([]); ws.append(cols); style_header(ws, 4, len(cols))
r = 5
for aspect, factor, val, why in ED:
    ws.cell(row=r, column=1, value=aspect); ws.cell(row=r, column=2, value=factor)
    ws.cell(row=r, column=3, value="0 / 2"); ws.cell(row=r, column=4, value=val)
    ws.cell(row=r, column=5, value=why)
    for c in range(1, len(cols) + 1):
        ws.cell(row=r, column=c).border = border; ws.cell(row=r, column=c).alignment = wrap
    ws.cell(row=r, column=3).alignment = ctr; ws.cell(row=r, column=4).alignment = ctr
    if val == 2:
        ws.cell(row=r, column=4).fill = PatternFill("solid", fgColor=LIGHT)
    r += 1
first, last = 5, r - 1
ws.cell(row=r, column=2, value="ED: TỔNG ĐIỂM MÔI TRƯỜNG").font = bold
ws.cell(row=r, column=4, value="=SUM(D%d:D%d)" % (first, last)).font = bold
ws.cell(row=r, column=4).fill = PatternFill("solid", fgColor=GREY); ws.cell(row=r, column=4).alignment = ctr
ws.cell(row=r, column=5, value="Tối đa 36")
r += 1
ws.cell(row=r, column=2, value="Hệ số môi trường = ED / 36").font = bold
ws.cell(row=r, column=4, value="=D%d/36" % (r - 1)).font = bold
ws.cell(row=r, column=4).number_format = "0.00"; ws.cell(row=r, column=4).alignment = ctr
r += 2
ws.cell(row=r, column=2, value="Tổng theo từng khía cạnh (tham khảo)").font = bold; r += 1
seen = []
for aspect, _, _, _ in ED:
    if aspect in seen:
        continue
    seen.append(aspect)
    rows = [i for i, e in enumerate(ED) if e[0] == aspect]
    ws.cell(row=r, column=2, value=aspect)
    ws.cell(row=r, column=4, value="=SUM(%s)" % ",".join("D%d" % (5 + i) for i in rows))
    ws.cell(row=r, column=4).alignment = ctr
    ws.cell(row=r, column=5, value="Tối đa 6")
    r += 1
widths(ws, [26, 74, 14, 13, 58])

# ============ 6. Ước lượng Story ============
ws = wb.create_sheet("6. Ước lượng Story")
ws["A1"] = "Ước lượng số lượng Story: UP / AP / PPS"; ws["A1"].font = title_font
ws["A2"] = "C (hệ số nhân) ="; ws["A2"].font = bold
ws["C2"] = C_DEFAULT; ws["C2"].fill = PatternFill("solid", fgColor=LIGHT); ws["C2"].font = bold
ws["D2"] = "ED ="; ws["D2"].font = bold
ws["E2"] = "='5. ED (Môi trường)'!D23"; ws["E2"].font = bold
ws["F2"] = "PPS = (AP × ED)/36 ;  AP = UP × C ;  UP = tổng 4 tiêu chí (mỗi tiêu chí 1..3)"
cols = ["ID", "User Story", "Loại tương tác\n(1-3)", "Quy tắc nghiệp vụ\n(1-3)", "Số thực thể\n(1-3)",
        "Thao tác dữ liệu\n(1-3)", "UP", "C", "AP", "ED", "PPS", "Story Point\n(Fibonacci)"]
ws.append([]); ws.append(cols); style_header(ws, 4, len(cols))
r = 5
for pb in PBS:
    ws.cell(row=r, column=1, value=pb["id"]); ws.cell(row=r, column=2, value=pb["user_story"])
    ws.cell(row=r, column=3, value=pb["tuong_tac"]); ws.cell(row=r, column=4, value=pb["quy_tac"])
    ws.cell(row=r, column=5, value=pb["thuc_the"]); ws.cell(row=r, column=6, value=pb["thao_tac"])
    ws.cell(row=r, column=7, value="=SUM(C%d:F%d)" % (r, r))
    ws.cell(row=r, column=8, value="=$C$2")
    ws.cell(row=r, column=9, value="=G%d*H%d" % (r, r))
    ws.cell(row=r, column=10, value="=$E$2")
    ws.cell(row=r, column=11, value="=I%d*J%d/36" % (r, r))
    ws.cell(row=r, column=12, value=("=IF(K{r}<=1,1,IF(K{r}<=2,2,IF(K{r}<=3,3,IF(K{r}<=5,5,"
                                     "IF(K{r}<=8,8,IF(K{r}<=13,13,21))))))").format(r=r))
    for c in range(1, len(cols) + 1):
        ws.cell(row=r, column=c).border = border
        ws.cell(row=r, column=c).alignment = ctr if c != 2 else wrap
    ws.cell(row=r, column=11).number_format = "0.00"
    r += 1
ws.cell(row=r, column=2, value="TỔNG").font = bold
ws.cell(row=r, column=7, value="=SUM(G5:G%d)" % (r - 1)).font = bold
ws.cell(row=r, column=11, value="=SUM(K5:K%d)" % (r - 1)).font = bold
ws.cell(row=r, column=11).number_format = "0.00"
ws.cell(row=r, column=12, value="=SUM(L5:L%d)" % (r - 1)).font = bold
widths(ws, [8, 52, 12, 13, 11, 12, 7, 7, 8, 7, 9, 12])

wb.save(OUT)
print("OK ->", OUT)
print("RQ =", NRQ, " PB =", NPB, " Task =", len(TASKS), " Tong gio =", sum(t["est"] for t in TASKS))
print("ED =", ED_TOTAL, "/", ED_MAX, "=> he so", round(ED_TOTAL / ED_MAX, 2))
